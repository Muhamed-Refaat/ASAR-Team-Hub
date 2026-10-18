"""Style the source workbook and add review-workflow lookups without moving source cells."""
from pathlib import Path
import json
import shutil
import subprocess
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as col

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'spreedsheet-copy/ASAR-Team-Hub.xlsx'
DELIVERY = ROOT / 'deliverables/ASAR-Team-Hub-Updated.xlsx'
NAVY, PURPLE, TEAL = '211C36', '5931B6', '087F8C'


def style(wb):
    from sync_local_db import upgrade
    upgrade(wb)
    # Rebuild only the generated reference sheets. Every original source cell stays in place.
    for name in ('Workbook Guide', 'Lookups'):
        if name in wb:
            del wb[name]
    guide = wb.create_sheet('Workbook Guide', 0)
    guide.append(['ASAR TEAM HUB', 'Workbook & review workflow'])
    notes = [
        ('START HERE', 'This workbook matches the shared frontend and Apps Script backend. Source cells are preserved; the inactive Exam sheet is hidden.'),
        ('Essential event fields', 'Tasks use title, description, assignees, status and due date. Meetings use title, date, participants, link/location and attendance logging. Courses use title, course link and assigned members, with optional track, requirement and deadline settings.'),
        ('Task workflow', 'Update task status using the Kanban drag-and-drop board or the status dropdown. There is no task completion form. Completed status counts for all assignees; moving back to an open status reverses that progress.'),
        ('Course workflow', 'Uncompleted > Pending > Completed. Evidence is optional. Members can submit without a file, or attach a PDF, PNG or JPEG up to 10 MB. Only an admin approval counts toward progress.'),
        ('Background saves', 'New events appear as saving drafts immediately while the app stays usable. A compact indicator beside Refresh shows pending changes and opens failure details for retry. Unconfirmed drafts do not count in progress. Retries reuse the same event ID to prevent duplicates.'),
        ('Event deletion', 'Only enabled admins can delete tasks, meetings and courses. Deleted TRUE on the source event and linked Log rows removes the event from assignments, progress and review queues. Completion attempts, evidence and Activity remain recorded. Queued announcements are cancelled. Use the app for consistent deletion; do not remove or reorder historical rows.'),
        ('References', 'Every active member can add a title, http/https URL and optional description in References. Only enabled admins can delete links. Added By, Added By Name, Added At and Deleted are app-managed fields; deletion preserves the source row and Activity history.'),
        ('Task evidence', 'Task assignees and admins add completion or issue links directly on the task card. Evidence Links is an app-managed field on Task that retains author, time, purpose, note and removed-link history. Adding a link never changes board status.'),
        ('Activity review', 'Activity includes user submissions, approvals, task evidence, settings changes and authenticated request failures. Use member/category/search filters and open details to investigate errors or request IDs. Completion retains the individual review records.'),
        ('Meeting workflow', 'An admin opens Logging Enabled. Members submit attendance for review; evidence is optional. Only approved attendance counts toward progress.'),
        ('Undo and update', 'Members can withdraw Pending or undo Completed records from their event card, including imported completions. Status returns to Uncompleted. Resubmission creates a new attempt and requires review again.'),
        ('Review', 'Admins use the review queue in Learning or Meetings. Approve sets Completed; Return sets Uncompleted and requires feedback. Withdrawn or superseded submissions cannot be approved.'),
        ('Completion sheet', 'Each row is a submission attempt. Submission ID is stable. Status, reviewer, review time, review note, last status change and actor are saved. Older attempts and attachments remain as history; the latest attempt is authoritative.'),
        ('Evidence attachments', 'Apps Script stores files in the submitting account\'s private ASAR evidence folder, shares each file with currently enabled admins, and writes its Drive ID, URL, name, type and size to Completion. Do not paste file bytes into cells.'),
        ('Historical records', 'Imported Log totals and existing Completion rows with blank Status remain accepted historical records. No reviewer or evidence is invented. A member undo overrides their imported completion.'),
        ('Review consistency', 'Use the web app to review or undo so Log aggregates and Activity history remain synchronized. Direct manual status edits are reflected by the app but do not create an authenticated review audit.'),
        ('Lookups and dropdowns', 'Lookups lists current members, admins, event keys, skills and allowed values. Completion Name, Email and Event Name contain lookup formulas for manual entry. Open in Excel or Google Sheets to calculate formulas.'),
        ('Manual data', 'Source event sheets define events. Invited/Assigned and Assignees accept exact names on separate lines. Team ID is the stable key. Admins is the only permission list; Role alone never grants admin access.'),
        ('Skills', 'The Skills page contains per-member required skills only. SkillRequirements uses Member ID, Skill, Category, TRUE/FALSE and optional Target Score (0-10). Without overrides, members inherit their recorded Team entries. Prior standalone assessments are archived; SkillEvents is hidden as historical data.'),
        ('Profile lists', 'Team Intelligence shows Technology, Skills, Knowledge, Abilities and Personality lists below the overview. Knowledge and Skills accept plain text lines or JSON arrays of category/items groups. Recorded Hard Skill and Soft Skill ratings also appear in Skills. Empty values stay empty.'),
        ('Legacy event columns', 'Older durations, task scores/feedback, track IDs, related skills, platform and meeting-type columns are preserved but hidden. New forms and saves use the essential fields only. Stable IDs, archive flags and review metadata remain operational fields.'),
        ('Subteams and roles', 'Subteams defines Subteam ID, Name, Description, Color and Active. TeamMemberships assigns Member ID and a Leader, Member or Support role per subteam. Each subteam needs one or more leaders; people can belong to multiple subteams. One overall Team Leader remains in Team.'),
        ('Team settings', 'Use Edit distribution in Team hierarchy to create/rename/remove subteams and assign all members. Global Leader roles, Department and Reports To are compatibility summaries; normalized membership rows store every relationship. Removed rows are disabled, never deleted.'),
        ('Member lifecycle', 'New members use the ASAR Team Application Google Form linked from the hub. The form-bound script in forms/TeamApplications.gs updates profile fields by header and preserves roles, permissions and removal state. Admins remove members in Team Intelligence: Team Active becomes FALSE and memberships are disabled; records remain. Transfer sole leadership first. Restore via the Team Active flag and Edit distribution when needed.'),
        ('Immediate updates', 'Task moves, completion submissions, reviews, edits, required skills and distribution changes update the UI immediately and save in the background. The compact header indicator shows pending changes. Failures roll back unconfirmed changes and keep Retry update available in its failure details. CV exports use confirmed data.'),
        ('Skill updates', 'Members can use Update result or Reset skill directly on their own skill card. Updates preserve prior results. Reset clears completion and removes that assessment from current ratings; original Team sheet ratings and history remain intact. Re-completion adds a fresh assessment.'),
        ('Dates', 'Meet Date is the scheduled date. Due to Date is the deadline through that date in Africa/Cairo. Members may undo even after a deadline; an admin must reopen a closed event before resubmission.'),
        ('Deployment', 'Import this workbook into Google Sheets, install generated Code.gs and Index.html, use Africa/Cairo and execute as the accessing signed-in user. Allow Drive authorization. Each user needs sheet write access. Adding a future admin may require sharing older evidence with them.'),
        ('Local preview', 'The local app keeps uploads as session-only files and provides download links. Local URLs are not persistent Drive evidence; deploy Apps Script for durable uploads. Export CSVs before reloading local data.'),
        ('Maintenance', 'Run python tools/style_workbook.py after changing the source schema or roster to refresh lookups and the delivery workbook. This also rebuilds CSVs and both app interfaces.'),
    ]
    for topic, text in notes:
        guide.append([topic, text])

    lookup = wb.create_sheet('Lookups')
    # Read the model so skill keys and synthetic event IDs exactly match the frontend.
    script = "const fs=require('fs'),M=require('./web/model.js');let db=JSON.parse(fs.readFileSync(0,'utf8'));process.stdout.write(JSON.stringify({events:M.events(db),skills:M.projectSkills(db),admins:db.team.filter(m=>M.isAdmin(db,m)),eventColumns:M.EVENT_COLUMNS}));"
    def records(name):
        sheet = wb[name]
        headers = [str(c.value).strip() if c.value is not None else '' for c in sheet[1]]
        return [{h: '' if values[i] is None else str(values[i]) for i,h in enumerate(headers) if h} for values in sheet.iter_rows(min_row=2, values_only=True) if any(v is not None for v in values)]
    db = {k: records(n) for k,n in {'team':'Team','admins':'Admins','tasks':'Task','meet':'Meet','learning':'Learning','logs':'Log','completions':'Completion','skillEvents':'SkillEvents','skillRequirements':'SkillRequirements','subteams':'Subteams','memberships':'TeamMemberships'}.items()}
    skill = wb['SKILLs'] if 'SKILLs' in wb else wb['Skill']
    skill_headers = ['Hard Skill','Technical Competency','Specific Technical Application','Soft Skill','Description','Engineering Layer','Technology','Function in Architecture']
    db['skills'] = [dict(zip(skill_headers, [str(v or '') for v in row[:8]])) for row in skill.iter_rows(min_row=3,values_only=True) if any(row)]
    model = json.loads(subprocess.run(['node','-e',script],input=json.dumps(db),text=True,capture_output=True,cwd=ROOT,check=True,encoding='utf-8').stdout)
    members = db['team']
    lists = {
        'Member ID': [m['ID'].removesuffix('.0') for m in members],
        'Member Name': [m['Name'] for m in members],
        'Member Email': [m['Mail'] for m in members],
        'Completion State': ['Uncompleted','Pending','Completed'],
        'Boolean': ['TRUE','FALSE'],
        'Role': ['Team member','Leader','Team Leader'],
        'Event Type': ['Task','Meet','Learning','Skill'],
        'Task Status': ['Backlog','Pending','InProgress','Blocked','Completed'],
        'Skill Category': ['Hard Skill','Soft Skill','Technology','Personality','Abilities'],
        'Skill': list(dict.fromkeys(s['label'] for s in model['skills'])),
        'Event Key': [e['key'] for e in model['events']],
        'Event Name': [e['title'] for e in model['events']],
        'Admin Email': [m['Mail'] for m in model['admins']],
        'Subteam ID': [t['Subteam ID'] for t in db['subteams'] if t['Active'].lower() not in ('false','0','no')],
        'Subteam Name': [t['Name'] for t in db['subteams'] if t['Active'].lower() not in ('false','0','no')],
        'Membership Role': ['Leader','Member','Support'],
        'Subteam Color': ['purple','teal','blue','amber','rose'],
    }
    lookup.append(list(lists))
    for row in range(max(map(len,lists.values()))):
        lookup.append([items[row] if row<len(items) else None for items in lists.values()])
    named = {'HubMemberIDs':'Member ID','HubMemberNames':'Member Name','HubMemberEmails':'Member Email','HubStates':'Completion State','HubBooleans':'Boolean','HubRoles':'Role','HubEventTypes':'Event Type','HubTaskStatuses':'Task Status','HubCategories':'Skill Category','HubSkillNames':'Skill','HubEventKeys':'Event Key','HubAdminEmails':'Admin Email','HubSubteamIDs':'Subteam ID','HubMembershipRoles':'Membership Role','HubTeamColors':'Subteam Color'}
    for name, field in named.items():
        n = list(lists).index(field)+1
        wb.defined_names[name] = DefinedName(name, attr_text=f"'Lookups'!${col(n)}$2:${col(n)}${max(2,len(lists[field])+1)}")
    wb.defined_names['HubProjectCategories'] = DefinedName('HubProjectCategories', attr_text="'Lookups'!$I$2:$I$4")

    validations = {
        'Team': {'Role':'HubRoles','Active':'HubBooleans'},
        'Subteams': {'Color':'HubTeamColors','Active':'HubBooleans'},
        'TeamMemberships': {'Subteam ID':'HubSubteamIDs','Member ID':'HubMemberIDs','Role':'HubMembershipRoles','Active':'HubBooleans'},
        'Admins': {'Email':'HubMemberEmails','Enabled':'HubBooleans'},
        'Task': {'Status':'HubTaskStatuses','Deleted':'HubBooleans'},
        'Meet': {'Logging Enabled':'HubBooleans','Archived':'HubBooleans','Deleted':'HubBooleans'},
        'Learning': {'Deleted':'HubBooleans'},
        'Log': {'Deleted':'HubBooleans'},
        'References': {'Deleted':'HubBooleans'},
        'SkillEvents': {'Skill Category':'HubCategories','Archived':'HubBooleans'},
        'SkillRequirements': {'Member ID':'HubMemberIDs','Skill':'HubSkillNames','Category':'HubProjectCategories','Required':'HubBooleans'},
        'Completion': {'Member ID':'HubMemberIDs','Status':'HubStates','Event Type':'HubEventTypes','Event ID':'HubEventKeys','Reviewer Email':'HubAdminEmails'},
    }
    for sheet in wb:
        is_skill = sheet.title in ('SKILLs','Skill')
        header = 2 if is_skill else 1
        sheet.freeze_panes = f'A{header+1}'
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1,0
        sheet.print_title_rows = f'1:{header}'
        sheet.sheet_properties.tabColor = PURPLE if sheet.title in ('Completion','Admins','Workbook Guide') else TEAL
        if sheet.title != 'Workbook Guide':
            sheet.auto_filter.ref = f'A{header}:{col(sheet.max_column)}{max(header+1,sheet.max_row)}'
        for row in sheet:
            for cell in row:
                is_header = cell.row <= header
                cell.font = Font(name='Aptos',size=11,color='FFFFFF' if is_header else '29243E',bold=is_header)
                cell.fill = PatternFill('solid',fgColor=NAVY if is_header else ('F4F1FC' if cell.row%2==0 else 'FFFFFF'))
                cell.alignment = Alignment(vertical='center',wrap_text=True)
                if isinstance(cell.value,str) and cell.value.startswith(('https://','http://')):
                    cell.hyperlink=cell.value
                    cell.font = Font(name='Aptos',size=11,color=PURPLE,underline='single')
            sheet.row_dimensions[row[0].row].height = 34 if row[0].row<=header else (70 if sheet.title=='Team' else 48)
        for c in range(1,sheet.max_column+1):
            label=str(sheet.cell(header,c).value or '')
            sheet.column_dimensions[col(c)].width = 25 if len(label)<22 else 34
            if label in ('Evidence','Review Note','Description','Task description','overview','Education','Knowledge','Skills'):
                sheet.column_dimensions[col(c)].width = 48
            if not label and all(sheet.cell(r,c).value is None for r in range(header+1,sheet.max_row+1)):
                sheet.column_dimensions[col(c)].hidden=True
            if sheet.title in model['eventColumns']:
                sheet.column_dimensions[col(c)].hidden = label not in model['eventColumns'][sheet.title]
        if sheet.title in validations:
            sheet.data_validations.dataValidation=[]
            headers={str(c.value).strip():c.column for c in sheet[header] if c.value}
            for label, defined in validations[sheet.title].items():
                if label not in headers or (sheet.title in model['eventColumns'] and label not in model['eventColumns'][sheet.title]): continue
                dv=DataValidation(type='list',formula1='='+defined,allow_blank=True)
                dv.errorTitle='Choose a listed value';dv.error='Use the dropdown list to keep the workbook aligned with the app.'
                dv.showErrorMessage=True;dv.showInputMessage=True;dv.promptTitle=label;dv.prompt='Select a value from Lookups.'
                sheet.add_data_validation(dv);dv.add(f'{col(headers[label])}2:{col(headers[label])}2000')
            for label in ('Score','Target Score'):
                if label not in headers: continue
                dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=10,allow_blank=True)
                dv.showErrorMessage=True;dv.error='Ratings must be between 0 and 10.'
                sheet.add_data_validation(dv);dv.add(f'{col(headers[label])}2:{col(headers[label])}2000')
        if sheet.title=='Completion':
            headers={c.value:c.column for c in sheet[1] if c.value}
            state=col(headers['Status'])
            sheet.conditional_formatting._cf_rules.clear()
            for status,bg,fg in [('Uncompleted','F0EDF5','5E566E'),('Pending','FFF0C7','805000'),('Completed','DCF5E7','176841')]:
                sheet.conditional_formatting.add(f'A2:{col(sheet.max_column)}2000',FormulaRule(formula=[f'${state}2="{status}"'],fill=PatternFill('solid',fgColor=bg),font=Font(color=fg)))
            # Blank-row formulas return empty strings, so importers and GAS ignore unused template rows.
            for r in range(2,302):
                for target,formula in {
                    'Name':f'IFERROR(INDEX(Lookups!$B$2:$B${len(members)+1},MATCH({col(headers["Member ID"])}{r},Lookups!$A$2:$A${len(members)+1},0)),"")',
                    'Email':f'IFERROR(INDEX(Lookups!$C$2:$C${len(members)+1},MATCH({col(headers["Member ID"])}{r},Lookups!$A$2:$A${len(members)+1},0)),"")',
                    'Event Name':f'IFERROR(INDEX(Lookups!$L$2:$L${len(model["events"])+1},MATCH({col(headers["Event ID"])}{r},Lookups!$K$2:$K${len(model["events"])+1},0)),"")',
                }.items():
                    cell=sheet.cell(r,headers[target])
                    if cell.value is None: cell.value='='+formula
            sheet.auto_filter.ref=f'A1:{col(sheet.max_column)}301'
        if sheet.title=='Meet':
            headers={c.value:c.column for c in sheet[1] if c.value}
            d=col(headers['Date']);sheet.conditional_formatting._cf_rules.clear()
            sheet.conditional_formatting.add(f'A2:{col(sheet.max_column)}2000',FormulaRule(formula=[f'AND(ISNUMBER(${d}2),${d}2>TODAY())'],fill=PatternFill('solid',fgColor='EEE5FF')))
    guide.column_dimensions['A'].width=27
    guide.column_dimensions['B'].width=112
    for r in range(2,guide.max_row+1): guide.row_dimensions[r].height=66
    guide.sheet_properties.tabColor=PURPLE
    if 'Exam' in wb: wb['Exam'].sheet_state='hidden'
    if 'SkillEvents' in wb: wb['SkillEvents'].sheet_state='hidden'
    wb.active=0
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)


def main():
    wb=openpyxl.load_workbook(BOOK)
    style(wb)
    wb.save(BOOK);wb.close()
    from sync_local_db import sync_database
    sync_database()
    DELIVERY.parent.mkdir(exist_ok=True)
    shutil.copy2(BOOK,DELIVERY)
    print(f'Delivery workbook: {DELIVERY}')

if __name__=='__main__':
    import sys
    sys.path.insert(0,str(ROOT))
    main()
