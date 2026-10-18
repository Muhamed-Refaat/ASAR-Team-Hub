"""Check source-cell preservation, CSV/snapshot parity, and real CV documents."""
import csv
import json
import re
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
from sync_local_db import serialize

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'spreedsheet-copy'

class DataAudit(unittest.TestCase):
    def test_delivery_workbook_matches_source_and_contains_review_schema_and_real_lookups(self):
        delivery = ROOT / 'deliverables/ASAR-Team-Hub-Updated.xlsx'
        self.assertEqual(delivery.read_bytes(), (DATA / 'ASAR-Team-Hub.xlsx').read_bytes())
        workbook = openpyxl.load_workbook(delivery)
        completion = workbook['Completion']
        headers = [c.value for c in completion[1]]
        for field in ('Submission ID', 'Status', 'Evidence File ID', 'Evidence URL', 'Evidence Filename', 'Evidence MIME Type', 'Evidence Size', 'Reviewer Email', 'Reviewer Name', 'Reviewed At', 'Review Note', 'Status Changed At', 'Status Changed By'):
            self.assertIn(field, headers)
        self.assertEqual(workbook.sheetnames[0], 'Workbook Guide')
        self.assertEqual(workbook['Exam'].sheet_state, 'hidden')
        self.assertIn('Logging Enabled', [c.value for c in workbook['Meet'][1]])
        self.assertEqual(completion.freeze_panes, 'A2')
        self.assertEqual(completion['A1'].fill.fgColor.rgb[-6:], '211C36')
        self.assertEqual(completion['A1'].font.color.rgb[-6:], 'FFFFFF')
        self.assertGreaterEqual(len(completion.conditional_formatting), 1)
        states = workbook.defined_names['HubStates']
        sheet, area = next(states.destinations)
        self.assertEqual([c.value for row in workbook[sheet][area] for c in row], ['Uncompleted', 'Pending', 'Completed'])
        ids = workbook.defined_names['HubMemberIDs']
        sheet, area = next(ids.destinations)
        self.assertEqual(len([c for row in workbook[sheet][area] for c in row if c.value]), 26)
        for sheet in ('Completion', 'Team', 'Admins', 'Meet', 'SkillRequirements', 'Subteams', 'TeamMemberships', 'References'):
            self.assertGreater(len(workbook[sheet].data_validations.dataValidation), 0, sheet)
        for field in ('Name', 'Email', 'Event Name'):
            cell = completion.cell(2, headers.index(field) + 1)
            self.assertEqual(cell.data_type, 'f')
            self.assertIn('Lookups!', cell.value)
            self.assertNotIn('#REF!', cell.value)
        for sheet in workbook:
            self.assertFalse(sheet.sheet_view.showGridLines)
        workbook.close()

    def test_subteams_match_member_ids_and_include_leaders_with_lookup_roles(self):
        workbook = openpyxl.load_workbook(DATA / 'ASAR-Team-Hub.xlsx', data_only=True)
        def rows(name):
            values = list(workbook[name].values)
            return [dict(zip(values[0], row)) for row in values[1:] if any(v is not None for v in row)]
        teams, memberships = rows('Subteams'), rows('TeamMemberships')
        ids = {serialize(m['ID']) for m in rows('Team') if m['Name']}
        self.assertEqual(len(teams), 3)
        self.assertEqual(len(memberships), 25)
        self.assertEqual(len({(m['Subteam ID'], m['Member ID']) for m in memberships}), len(memberships))
        for row in memberships:
            self.assertIn(serialize(row['Member ID']), ids)
            self.assertIn(row['Role'], ('Leader', 'Member', 'Support'))
        for team in teams:
            self.assertTrue(any(m['Subteam ID'] == team['Subteam ID'] and m['Role'] == 'Leader' for m in memberships))
        sheet, area = next(workbook.defined_names['HubMembershipRoles'].destinations)
        self.assertEqual([c.value for row in workbook[sheet][area] for c in row], ['Leader', 'Member', 'Support'])
        workbook.close()

    def test_event_sheets_show_essential_fields_and_preserve_hidden_legacy_columns(self):
        workbook = openpyxl.load_workbook(DATA / 'ASAR-Team-Hub.xlsx')
        from openpyxl.utils import get_column_letter
        visible = lambda name: {c.value for c in workbook[name][1] if c.value and not workbook[name].column_dimensions[get_column_letter(c.column)].hidden}
        self.assertEqual(visible('Task'), {'Task Title', 'Task description', 'Assignees', 'Status', 'Due to Date', 'Archived', 'Deleted', 'Event ID'})
        self.assertEqual(visible('Meet'), {'Meeting title', 'Date', 'Description', 'Attachments', 'Invited/Assigned', 'Logging Enabled', 'Due to Date', 'Archived', 'Deleted', 'Event ID'})
        self.assertEqual(visible('Learning'), {'Topic name', 'Link', 'Track Name', 'Obligability', 'Invited/Assigned', 'Due to Date', 'Description', 'Deleted', 'Event ID'})
        self.assertIn('Deleted', [c.value for c in workbook['Log'][1]])
        self.assertEqual(workbook['SkillEvents'].sheet_state, 'hidden')
        self.assertIn('Knowledge', [c.value for c in workbook['Team'][1]])
        self.assertIn('Skills', [c.value for c in workbook['Team'][1]])
        self.assertIn('Evidence Links', [c.value for c in workbook['Task'][1]])
        self.assertIn('Active', [c.value for c in workbook['Team'][1]])
        guide = dict(workbook['Workbook Guide'].iter_rows(min_row=2, values_only=True))
        self.assertIn('Evidence is optional', guide['Course workflow'])
        self.assertIn('saving drafts immediately', guide['Background saves'])
        self.assertIn('forms/TeamApplications.gs', guide['Member lifecycle'])
        self.assertIn('Completion attempts, evidence and Activity remain recorded', guide['Event deletion'])
        self.assertIn('Every active member can add', guide['References'])
        self.assertEqual([c.value for c in workbook['References'][1]], ['Reference ID', 'Title', 'URL', 'Description', 'Added By', 'Added By Name', 'Added At', 'Deleted'])
        workbook.close()

    def test_original_nonempty_cells_are_preserved(self):
        if not (DATA / 'ASAR-Team-Hub.before-schema.xlsx').exists():
            self.skipTest('The initial workbook backup is kept locally during import.')
        before = openpyxl.load_workbook(DATA / 'ASAR-Team-Hub.before-schema.xlsx', data_only=True, read_only=True)
        after = openpyxl.load_workbook(DATA / 'ASAR-Team-Hub.xlsx', data_only=True, read_only=True)
        for sheet in before:
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        self.assertEqual(cell.value, after[sheet.title][cell.coordinate].value, f'{sheet.title}!{cell.coordinate}')
        before.close()
        after.close()

    def test_csv_and_embedded_snapshot_are_identical(self):
        text = (ROOT / 'local_test_dashboard.html').read_text(encoding='utf-8')
        seed = json.loads(re.search(r'const LOCAL_SEED = (.*);\n', text).group(1))
        for key, sheet in seed['sheets'].items():
            self.assertEqual(seed['csv'][key], (DATA / f'{sheet}.csv').read_text(encoding='utf-8'))
        self.assertEqual(len(list(csv.DictReader(seed['csv']['team'].splitlines()))), 26)
        self.assertEqual(len(list(csv.DictReader(seed['csv']['activity'].splitlines()))), 35)
        self.assertEqual(len(list(csv.DictReader(seed['csv']['skills'].splitlines()))), 11)

    def test_workbook_rows_match_csv_rows_including_sparse_skill_headers(self):
        workbook = openpyxl.load_workbook(DATA / 'ASAR-Team-Hub.xlsx', data_only=True)
        for sheet in workbook:
            name = 'Skill' if sheet.title.lower() in ('skill', 'skills') else sheet.title
            header_row = 2 if name == 'Skill' else 1
            columns = list(range(1,9)) if name == 'Skill' else [c.column for c in sheet[1] if c.value is not None and str(c.value).strip()]
            expected = [[serialize(sheet.cell(r,c).value) for c in columns] for r in range(header_row+1,sheet.max_row+1) if any(sheet.cell(r,c).value is not None and str(sheet.cell(r,c).value).strip() for c in columns)]
            with (DATA / f'{name}.csv').open(encoding='utf-8',newline='') as handle:
                records = list(csv.reader(handle))
            self.assertEqual(records[1:],expected,name)
        workbook.close()

    def test_generated_interfaces_share_model_theme_and_react_app(self):
        for name in ('Index.html','local_test_dashboard.html'):
            content = (ROOT/name).read_text(encoding='utf-8')
            for source in ('model.js','cv.js','sync.js','actions.js','telegram.js','theme.css','app.jsx'):
                self.assertIn((ROOT/'web'/source).read_text(encoding='utf-8'),content)
        self.assertIn((ROOT/'web/model.js').read_text(encoding='utf-8'),(ROOT/'Code.gs').read_text(encoding='utf-8'))
        self.assertIn((ROOT/'web/telegram.js').read_text(encoding='utf-8'),(ROOT/'Code.gs').read_text(encoding='utf-8'))
        self.assertNotIn('LOCAL_SEED',(ROOT/'Index.html').read_text(encoding='utf-8'))

    def test_actual_word_exports_have_professional_layout_complete_skills_and_arabic_direction(self):
        exports = json.loads((ROOT/'test-results/cv-exports.json').read_text(encoding='utf-8'))
        ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for index, entry in enumerate(exports):
            with zipfile.ZipFile(ROOT/'test-results'/entry['file']) as archive:
                document = ET.fromstring(archive.read('word/document.xml'))
                text = '\n'.join(t.text or '' for t in document.findall('.//w:t',ns))
                self.assertIn(entry['name'],text)
                self.assertIn('ENGINEERING SKILLS',text)
                self.assertIn('TECHNICAL TOOLKIT',text)
                self.assertNotIn('TEAM PARTICIPATION',text)
                self.assertNotIn('/10',text)
                self.assertNotIn('undefined',text)
                self.assertNotIn('null',text)
                for label in entry['technical'] + entry['tools']:
                    self.assertIn(label,text)
                self.assertIn('Gazebo / RViz2 / WebSockets Injection',text)
                self.assertNotIn('Bachelor',text)
                self.assertNotIn('Expert',text)
                page = document.find('.//w:pgSz',ns)
                self.assertEqual(page.attrib['{'+ns['w']+'}w'],'11906')
                self.assertEqual(page.attrib['{'+ns['w']+'}h'],'16838')
                self.assertTrue(document.findall('.//w:keepNext',ns))
                self.assertTrue(document.findall('.//w:tbl',ns))
                self.assertTrue(document.findall('.//w:shd',ns))
                self.assertTrue(document.findall('.//w:hyperlink',ns))
                footer = ET.fromstring(archive.read('word/footer1.xml'))
                self.assertIn('NUMPAGES',' '.join(t.text or '' for t in footer.findall('.//w:instrText',ns)))
                if index == 1:
                    self.assertTrue(document.findall('.//w:bidi',ns))
                    self.assertTrue(document.findall('.//w:rtl',ns))

if __name__ == '__main__':
    unittest.main()
